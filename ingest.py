"""Excel to canonical rows.

Maps messy partner header casing/wording to canonical field names via a
hardcoded synonym dictionary. Trims whitespace on all values. Case
canonicalisation of `size_unit` is deferred to the schema layer, which
knows the row's category and can match units case-insensitively to that
category's allowed set (avoids the "L" ambiguity between liters and
apparel size L).

Unknown columns are dropped with a warning returned alongside.
"""

from typing import Any

from openpyxl import load_workbook

# Canonical field -> set of accepted synonyms (all normalised to lowercase, stripped).
HEADER_SYNONYMS: dict[str, set[str]] = {
    "category":     {"category", "cat", "product category"},
    "brand":        {"brand", "brand name", "manufacturer"},
    "title":        {"title", "product title", "product name", "name"},
    "gtin":         {"gtin", "ean", "upc", "ean/gtin", "barcode", "ean gtin"},
    "size":         {"size", "product size", "size value"},
    "size_unit":    {"size unit", "unit", "uom", "size_unit"},
    "image_url":    {"image url", "image link", "image", "main image", "image_url"},
    "video_url":    {"video url", "video link", "video", "video_url"},
    "description":  {"description", "desc", "product description", "long description"},
    "partner_sku":  {"partner sku", "sku", "seller sku", "vendor sku", "partner_sku"},
}


def _normalise_header(h: Any) -> str:
    if h is None:
        return ""
    return " ".join(str(h).lower().strip().split())


def _build_header_map(headers: list[Any]) -> tuple[dict[int, str], list[str]]:
    """Return (col_index -> canonical_field, list_of_unknown_headers)."""
    mapping: dict[int, str] = {}
    unknown: list[str] = []
    for idx, raw in enumerate(headers):
        norm = _normalise_header(raw)
        if not norm:
            continue
        found = None
        for canonical, syns in HEADER_SYNONYMS.items():
            if norm in syns:
                found = canonical
                break
        if found:
            mapping[idx] = found
        else:
            unknown.append(str(raw))
    return mapping, unknown


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def ingest(path: str) -> list[dict]:
    """Read the workbook and return a list of canonical Row dicts.

    Row shape:
        {
          "row_id": "r001",
          "fields": { canonical field -> str },
          "flags": [],
          "verdict_schema": None,
          "verdict_semantic": None,
        }
    """
    rows, _ = ingest_with_warnings(path)
    return rows


def ingest_with_warnings(path: str) -> tuple[list[dict], list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    warnings: list[str] = []

    iter_rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(iter_rows)
    except StopIteration:
        return [], ["empty worksheet"]

    header_map, unknown = _build_header_map(list(header_row))
    for h in unknown:
        warnings.append(f"unknown column ignored: {h!r}")

    canonical_fields = list(HEADER_SYNONYMS.keys())
    rows: list[dict] = []
    for i, raw_row in enumerate(iter_rows):
        # Skip fully blank lines.
        if raw_row is None or all(v is None or str(v).strip() == "" for v in raw_row):
            continue
        fields = {k: "" for k in canonical_fields}
        for col_idx, canonical in header_map.items():
            if col_idx < len(raw_row):
                fields[canonical] = _cell_to_str(raw_row[col_idx])
        rows.append({
            "row_id": f"r{i + 1:03d}",
            "fields": fields,
            "flags": [],
            "verdict_schema": None,
            "verdict_semantic": None,
        })
    return rows, warnings
