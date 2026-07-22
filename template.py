"""Template ingestion — a marketplace's own Excel template becomes a rulepack.

"Onboard the AI employee the way you'd onboard a human: hand it the
template." No per-customer configuration project. Signals we extract from
the .xlsx:

- Header row 1  → the field list.
- Any coloured header cell (non-white, non-empty fill) → that field is
  required. This mirrors the "required columns are highlighted" convention
  most marketplaces use in their own template documentation.
- A `list`-type data validation on a column → the enum for that field.
- (Optional) A sheet named `Config` with `category | size_units`
  (comma-separated) rows → per-category size_unit rules.

We never touch user data; only row-1 fills and workbook-level data
validations. The generated rulepack is saved per-tenant so different
customers can bring different templates without stepping on each other.
"""

import json
import os
from datetime import datetime, timezone

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------

def _rgb_is_meaningful(rgb) -> bool:
    """A header fill counts as 'required' if it's a real colour, not the
    workbook default of transparent/white."""
    if not rgb:
        return False
    s = str(rgb).upper()
    if len(s) < 6:
        return False
    return s not in {"00000000", "FFFFFFFF", "FFFFFFFFFF"}


def _normalise_header(raw) -> str:
    """Turn 'Product Title', 'EAN/GTIN', 'Size Unit ' into canonical names."""
    if raw is None:
        return ""
    s = str(raw).strip().lower()
    # A tiny synonym pass so partner-style header wording lands on our field ids.
    synonyms = {
        "brand name": "brand",
        "product title": "title",
        "ean/gtin": "gtin",
        "ean": "gtin",
        "size unit": "size_unit",
        "image link": "image_url",
        "video link": "video_url",
        "seller sku": "partner_sku",
    }
    if s in synonyms:
        return synonyms[s]
    return s.replace(" ", "_").replace("/", "_")


def _extract_headers(ws):
    """Return (col_index → field_name, field_name → required?)."""
    headers: dict[int, str] = {}
    required: dict[str, bool] = {}
    for cell in ws[1]:
        if cell.value is None:
            continue
        name = _normalise_header(cell.value)
        if not name:
            continue
        headers[cell.column] = name
        is_required = False
        try:
            fill = cell.fill
            if fill and fill.fgColor and _rgb_is_meaningful(fill.fgColor.rgb):
                is_required = True
        except Exception:
            pass
        required[name] = is_required
    return headers, required


def _parse_list_formula(formula: str | None, wb) -> list[str]:
    """Data validations of type 'list' store either an inline quoted list
    (`"ml,g,kg"`) or a range reference (`Config!$A$2:$A$10`). Handle both."""
    if not formula:
        return []
    f = str(formula).strip().lstrip("=")
    if f.startswith('"') and f.endswith('"'):
        return [v.strip() for v in f[1:-1].split(",") if v.strip()]
    if "!" in f:
        sheet_name, addr = f.split("!", 1)
        sheet_name = sheet_name.strip("'")
    else:
        sheet_name = wb.sheetnames[0]
        addr = f
    addr = addr.replace("$", "")
    if sheet_name not in wb.sheetnames:
        return []
    ref = wb[sheet_name]
    out: list[str] = []
    try:
        if ":" in addr:
            for row in ref[addr]:
                for cell in row:
                    if cell.value is not None:
                        out.append(str(cell.value).strip())
        else:
            v = ref[addr].value
            if v is not None:
                out.append(str(v).strip())
    except Exception:
        return []
    # De-dup preserving order.
    seen, result = set(), []
    for v in out:
        if v and v not in seen:
            seen.add(v)
            result.append(v)
    return result


def _extract_enums(ws, wb, headers) -> dict[str, list[str]]:
    enums: dict[str, list[str]] = {}
    dvs = getattr(ws.data_validations, "dataValidation", []) or []
    for dv in dvs:
        if getattr(dv, "type", None) != "list":
            continue
        allowed = _parse_list_formula(dv.formula1, wb)
        if not allowed:
            continue
        sqref = dv.sqref
        ranges = list(sqref.ranges) if hasattr(sqref, "ranges") else [sqref]
        for r in ranges:
            min_col = getattr(r, "min_col", None)
            max_col = getattr(r, "max_col", None) or min_col
            if min_col is None:
                try:
                    first = str(r).split(":")[0]
                    col_letter, _row = coordinate_from_string(first)
                    min_col = column_index_from_string(col_letter)
                    max_col = min_col
                except Exception:
                    continue
            for c in range(min_col, max_col + 1):
                if c in headers:
                    enums[headers[c]] = allowed
    return enums


def _extract_categories(wb) -> dict | None:
    """Optional Config sheet: row 1 is header, subsequent rows are
    (category, size_units) — units comma-separated in the second column."""
    if "Config" not in wb.sheetnames:
        return None
    ws = wb["Config"]
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        cat = str(row[0]).strip()
        units = row[1] if len(row) > 1 else None
        unit_list = [u.strip() for u in str(units or "").split(",") if u.strip()]
        out[cat] = {"size_units": unit_list, "size_range": None}
    return out or None


def read_template(path: str) -> dict:
    """Parse an .xlsx template and return a rulepack dict."""
    wb = load_workbook(path, data_only=False)
    ws = wb.active
    headers, required = _extract_headers(ws)
    enums = _extract_enums(ws, wb, headers)
    categories = _extract_categories(wb)

    fields = []
    for col_idx in sorted(headers.keys()):
        name = headers[col_idx]
        fields.append({
            "name": name,
            "required": required.get(name, False),
            "enum": enums.get(name),
        })

    return {
        "source": os.path.basename(path),
        "generated_at": datetime.now(timezone.utc)
                                .isoformat(timespec="seconds").replace("+00:00", "Z"),
        "fields": fields,
        "categories": categories,
    }


# ---------------------------------------------------------------------------
# Persistence — one rulepack per tenant, sits next to their corrections.
# ---------------------------------------------------------------------------

_HERE = os.path.dirname(os.path.abspath(__file__))
_TENANTS_DIR = os.path.join(_HERE, "data", "tenants")


def _rulepack_path(tenant: str) -> str:
    return os.path.join(_TENANTS_DIR, tenant, "rulepack.json")


def save_rulepack(rulepack: dict, tenant: str) -> None:
    path = _rulepack_path(tenant)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(rulepack, f, indent=2)


def load_rulepack(tenant: str) -> dict | None:
    path = _rulepack_path(tenant)
    if not os.path.exists(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def clear_rulepack(tenant: str) -> None:
    path = _rulepack_path(tenant)
    if os.path.exists(path):
        os.remove(path)
