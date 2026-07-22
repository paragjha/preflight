"""Error-sheet ingestion — the marketplace's returned error xlsx becomes memory.

When the cataloging tool rejects rows, its error sheet is labelled training
data: for each rejected row it names the field that broke and the reason.
We fold that back into the tenant's `corrections.json` as `confirmed_error`
records so the next semantic run treats similar rows as suspect — the
"the system that rejects you becomes the system that trains you" loop.

Expected columns (case-insensitive, best-effort matching):

    Match by (first present wins):
        row_id | rowid | id
        partner_sku | sku | seller_sku
        gtin | ean | ean/gtin

    Diagnostic:
        field | column | attribute — the field the tool blamed
        error | error_message | reason | message — the tool's message

If the sheet has none of the match columns, we return no corrections rather
than invent them.
"""

from openpyxl import load_workbook


_MATCH_KEYS_ID   = {"row_id", "rowid", "id"}
_MATCH_KEYS_SKU  = {"partner_sku", "sku", "seller_sku"}
_MATCH_KEYS_GTIN = {"gtin", "ean", "ean/gtin"}
_FIELD_KEYS      = {"field", "column", "attribute"}
_REASON_KEYS     = {"error", "error_message", "reason", "message"}


def _norm(s) -> str:
    if s is None:
        return ""
    return str(s).strip().lower().replace(" ", "_")


def parse_error_sheet(path: str) -> list[dict]:
    """Return a list of error records: {'match': {...}, 'field': str, 'reason': str}."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return []
    headers = [_norm(h) for h in header_row]
    col = {h: i for i, h in enumerate(headers) if h}

    def _pick(keys):
        for k in keys:
            if k in col:
                return col[k]
        return None

    id_i    = _pick(_MATCH_KEYS_ID)
    sku_i   = _pick(_MATCH_KEYS_SKU)
    gtin_i  = _pick(_MATCH_KEYS_GTIN)
    field_i = _pick(_FIELD_KEYS)
    reason_i = _pick(_REASON_KEYS)

    if id_i is None and sku_i is None and gtin_i is None:
        return []

    out: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        match: dict[str, str] = {}
        if id_i   is not None and id_i   < len(row) and row[id_i]   is not None:
            match["row_id"]      = str(row[id_i]).strip()
        if sku_i  is not None and sku_i  < len(row) and row[sku_i]  is not None:
            match["partner_sku"] = str(row[sku_i]).strip()
        if gtin_i is not None and gtin_i < len(row) and row[gtin_i] is not None:
            match["gtin"]        = str(row[gtin_i]).strip()
        if not match:
            continue
        out.append({
            "match":  match,
            "field":  str(row[field_i]).strip()  if field_i  is not None and field_i  < len(row) and row[field_i]  is not None else "",
            "reason": str(row[reason_i]).strip() if reason_i is not None and reason_i < len(row) and row[reason_i] is not None else "",
        })
    return out


def resolve_to_rows(error_records: list[dict], batch: dict) -> list[dict]:
    """Attach the matching batch row to each error record.

    Returns records with an added 'row' key. Records that couldn't be
    matched (no batch row found) get row=None so the caller can count
    unmatched separately.
    """
    rows = batch.get("rows", [])
    by_row_id = {r.get("row_id"): r for r in rows}
    by_sku    = {r["fields"].get("partner_sku"): r for r in rows
                 if r["fields"].get("partner_sku")}
    by_gtin   = {r["fields"].get("gtin"): r for r in rows
                 if r["fields"].get("gtin")}

    resolved = []
    for rec in error_records:
        m = rec["match"]
        r = None
        if m.get("row_id"):
            r = by_row_id.get(m["row_id"])
        if r is None and m.get("partner_sku"):
            r = by_sku.get(m["partner_sku"])
        if r is None and m.get("gtin"):
            r = by_gtin.get(m["gtin"])
        resolved.append({**rec, "row": r})
    return resolved


def as_corrections(resolved: list[dict]) -> list[dict]:
    """Turn resolved error records into memory-compatible correction dicts."""
    out = []
    for rec in resolved:
        row = rec.get("row")
        if row is None:
            continue
        out.append({
            "fields": row["fields"],
            "human_verdict": "confirmed_error",
            "note": f"validator returned: {rec['reason']}".strip(": ") if rec["reason"]
                    else "returned by the marketplace validator",
            "flag_code": rec["field"] or "OTHER",
        })
    return out
